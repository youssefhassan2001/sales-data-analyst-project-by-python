import pandas as pd
import matplotlib.pyplot as plt

df = pd.read_excel(
    r"C:\Users\JOOO\PycharmProjects\SQL\sales_training.xlsx",
    engine="openpyxl"
)

print(df.head())
print(df.shape)
df["Sales"]=df["Sales"].fillna(df["Sales"].mean())
df["Profit"]=df["Profit"].fillna(df["Profit"].mean())
print(df.isna().sum())
df=df.drop_duplicates()
print(df.duplicated())
print(df.dtypes)
df["Customer"]=df["Customer"].str.strip()
df["City"]=df["City"].str.upper()
df["Category"]=df["Category"].str.upper()
totalSales=df["Sales"].sum()
totalprofit=df["Profit"].sum()
print(f"your total sales is :{totalSales:,.2f} and achived {totalprofit:,.2f} as total profit")
topsales=df.groupby("Product")["Sales"].sum()
print(topsales.idxmax(),f" and achived {topsales.max():,.2f}")
topcity=df.groupby("City")["Sales"].sum()
print(topcity.idxmax(),f" and achived {topcity.max():,.2f}")
topCategory=df.groupby("Category")["Sales"].sum()
print(topCategory.idxmax(),f" and achived {topCategory.max():,.2f}")
pofitMargin=(totalprofit/totalSales)*100
print(type(topsales))
print(f"your profit margin is :{pofitMargin:.2f}%")
num=df["Quantity"].sum()
print(f"the total Quantity sold is : {num}")
Average_Order_Value =totalSales/num
total_orders = df["Order_ID"].nunique()
average_order_value = totalSales / total_orders
print(f"Average Order Value: {average_order_value:,.2f}")
plt.xlabel("Product")
topsal = topsales.sort_values(ascending=False)
print(topsal)
plt.bar(topsal.index,topsal.values)
plt.ylabel("Sales")
plt.title("Sales by product")
plt.grid(axis="y",alpha=.3)
for i ,value in enumerate(topsal.values):
    plt.text(i,value,str(value),ha="center")
plt.show()
topcity = topcity.sort_values(ascending=False)
plt.bar(topcity.index,topcity.values)
plt.ylabel("Sales")
plt.title("Sales by city")
plt.grid(axis="y",alpha=.3)
for i ,value in enumerate(topcity.values):
    plt.text(i,value,str(value),ha="center")
plt.show()
topCategory = topCategory.sort_values(ascending=False)
plt.bar(topCategory.index,topCategory.values)
plt.ylabel("Sales")
plt.title("Sales by Category")
plt.grid(axis="y",alpha=.3)
for i ,value in enumerate(topCategory.values):
    plt.text(i,value,f"{value}",ha="center")
plt.show()
df["Month"] = df["Date"].dt.month

monthly_sales = df.groupby("Month")["Sales"].sum()

plt.plot(monthly_sales.index, monthly_sales.values, marker="o")

plt.ylabel("Sales")
plt.xlabel("Month")
plt.title("Sales by Month")
plt.grid(axis="y", alpha=.3)
for month, value in zip(monthly_sales.index, monthly_sales.values):
    plt.text(month, value, f"{value:,.0f}", ha="center")
plt.show()
profit_product = df.groupby("Product")["Profit"].sum()
profit_product = profit_product.sort_values(ascending=False)
plt.bar(profit_product.index, profit_product.values)

plt.ylabel("profit")
plt.xlabel("product")
plt.title("profit by product")
plt.grid(axis="y", alpha=.3)
for i, value in zip(profit_product.index, profit_product.values):
    plt.text(i, value, f"{value:,.0f}", ha="center")
plt.show()
product_analysis = df.groupby("Product")[["Sales", "Profit"]].sum()
product_analysis["Profit_Margin"] = (
    product_analysis["Profit"] / product_analysis["Sales"]
) * 100
product_analysis = product_analysis.sort_values(
    "Profit_Margin",
    ascending=False
)

print(product_analysis)
best_margin_product = product_analysis["Profit_Margin"].idxmax()

print(best_margin_product)
best_sales_product = topsales.idxmax()
best_sales = topsales.max()

print(f"Top Product: {best_sales_product}")
print(f"Sales: {best_sales:,.2f}")
laptop_sales_percentage = (best_sales / totalSales) * 100

print(f"Laptop Sales Percentage: {laptop_sales_percentage:.2f}%")
best_city = topcity.idxmax()
best_city_sales = topcity.max()

city_percentage = (best_city_sales / totalSales) * 100

print(f"Top City: {best_city}")
print(f"Sales Percentage: {city_percentage:.2f}%")
summary = pd.DataFrame({
    "Metric": [
        "Total Sales",
        "Total Profit",
        "Profit Margin",
        "Total Quantity Sold",
        "Average Order Value",
        "Top Product",
        "Top City"
    ],

    "Value": [
        totalSales,
        totalprofit,
        f"{pofitMargin:.2f}%",
        num,
        average_order_value,
        best_sales_product,
        best_city
    ]
})

print(summary)
with pd.ExcelWriter("sales_analysis.xlsx", engine="openpyxl") as writer:

    # 1️⃣ KPI Summary
    summary.to_excel(
        writer,
        sheet_name="KPI Summary",
        index=False
    )

    # 2️⃣ Product Analysis
    product_analysis.to_excel(
        writer,
        sheet_name="Product Analysis"
    )

    # 3️⃣ City Sales
    topcity.to_frame("Sales").to_excel(
        writer,
        sheet_name="City Sales"
    )

    # 4️⃣ Category Sales
    topCategory.to_frame("Sales").to_excel(
        writer,
        sheet_name="Category Sales"
    )

    # 5️⃣ Monthly Sales
    monthly_sales.to_frame("Sales").to_excel(
        writer,
        sheet_name="Monthly Sales"
    )

    # 6️⃣ Profit by Product
    profit_product.to_frame("Profit").to_excel(
        writer,
        sheet_name="Profit by Product"
    )

print("Excel file created successfully!")
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

file_name = "sales_analysis.xlsx"

# فتح الملف
wb = load_workbook(file_name)

for ws in wb.worksheets:

    # تنسيق Header
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")

    # ضبط عرض الأعمدة تلقائيًا
    for column in ws.columns:
        max_length = 0
        column_letter = get_column_letter(column[0].column)

        for cell in column:
            if cell.value is not None:
                max_length = max(max_length, len(str(cell.value)))

        ws.column_dimensions[column_letter].width = max_length + 3

    # Freeze Header
    ws.freeze_panes = "A2"

    # Filter
    ws.auto_filter.ref = ws.dimensions


# تنسيق KPI Summary
ws = wb["KPI Summary"]

for row in range(2, ws.max_row + 1):

    metric = ws.cell(row=row, column=1).value
    value_cell = ws.cell(row=row, column=2)

    # Bold Metrics
    ws.cell(row=row, column=1).font = Font(bold=True)

    # تنسيق الأرقام
    if metric in ["Total Sales", "Total Profit", "Average Order Value"]:
        value_cell.number_format = '#,##0.00'

    elif metric == "Total Quantity Sold":
        value_cell.number_format = '#,##0'


# حفظ الملف
wb.save(file_name)

print("Professional Excel report created successfully!")